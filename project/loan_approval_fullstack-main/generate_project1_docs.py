"""
PROJECT 1: Loan Approval Prediction System
Generates:
  - Loan_Approval_Report.docx
  - Loan_Approval_Presentation.pptx
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

from pptx import Presentation
from pptx.util import Inches as PI, Pt as PPt
from pptx.dml.color import RGBColor as PRC
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DOCX_OUT = os.path.join(BASE_DIR, "Loan_Approval_Report.docx")
PPTX_OUT = os.path.join(BASE_DIR, "Loan_Approval_Presentation.pptx")

def shade_cell(cell, hex_color):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))

def cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val)); node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)

def doc_h(doc, text, level, rgb=(31, 78, 120)):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(*rgb)
    return h

def build_docx():
    print("Generating Loan_Approval_Report.docx ...")
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin = sec.right_margin = Inches(1)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LOAN APPROVAL PREDICTION SYSTEM")
    r.font.name, r.font.size, r.font.bold = "Arial", Pt(24), True
    r.font.color.rgb = RGBColor(31, 78, 120)

    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Machine Learning Classification Report  |  FastAPI REST Backend  +  Streamlit Web UI\n")
    r2.font.name, r2.font.size, r2.font.italic = "Arial", Pt(13), True
    r2.font.color.rgb = RGBColor(89, 89, 89)

    doc_h(doc, "1. Executive Summary", 1)
    doc.add_paragraph(
        "This report documents the end-to-end Loan Approval Prediction System: a machine learning "
        "pipeline (loan.py) trained on 4,269 customer records, deployed via a FastAPI REST backend "
        "(backend/app.py) and accessible through an interactive Streamlit web dashboard "
        "(frontend/streamlit_app.py). The system returns real-time continuous probability scores "
        "for loan approval decisions."
    )
    tbl = doc.add_table(1, 1); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0); shade_cell(c, "EBF3FB"); cell_margins(c, 150, 150, 200, 200)
    bp = c.paragraphs[0]
    r_k = bp.add_run("KEY RESULT: "); r_k.font.bold, r_k.font.color.rgb = True, RGBColor(31, 78, 120)
    bp.add_run("Random Forest Classifier achieved ")
    for t in ["97.89% Test Accuracy", " · ", "97.61% Precision", " · ", "99.07% Recall", " · ", "98.33% F1-Score"]:
        rb = bp.add_run(t); rb.font.bold = (t != " · ")

    doc_h(doc, "2. Problem Definition", 1)
    doc.add_paragraph("Financial institutions face credit risk from loan defaults. Manual underwriting is slow and inconsistent. An automated ML system evaluates applications consistently at scale.")
    for item in [
        "Business Goal: Reduce Non-Performing Assets (NPAs) by flagging high-risk applications early.",
        "Target Variable: loan_status — binary classification (Approved=1, Rejected=0).",
        "Dataset: loan_approval_dataset.csv — 4,269 customer records, zero missing values.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc_h(doc, "3. Dataset Overview & Feature Engineering", 1)
    doc_h(doc, "3.1 Raw Features", 2)
    feat_tbl = doc.add_table(9, 3); feat_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ci, h in enumerate(["Feature", "Type", "Description"]):
        c = feat_tbl.rows[0].cells[ci]; c.text = h
        shade_cell(c, "1F4E78")
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    rows = [
        ("no_of_dependents", "Integer", "Number of financial dependents"),
        ("education", "Categorical", "Graduate / Not Graduate"),
        ("self_employed", "Categorical", "Yes / No"),
        ("income_annum", "Integer (INR)", "Annual income of applicant"),
        ("loan_amount", "Integer (INR)", "Requested loan principal"),
        ("loan_term", "Integer (months)", "Requested repayment duration"),
        ("cibil_score", "Integer (300–900)", "Credit bureau score — primary risk indicator"),
        ("total_assets_value", "Engineered", "residential + commercial + luxury + bank assets"),
    ]
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row):
            feat_tbl.rows[ri].cells[ci].text = val
            cell_margins(feat_tbl.rows[ri].cells[ci])

    doc_h(doc, "3.2 Asset Aggregation Formula", 2)
    p_formula = doc.add_paragraph()
    r_f = p_formula.add_run("total_assets_value = residential_assets_value + commercial_assets_value + luxury_assets_value + bank_asset_value")
    r_f.font.name, r_f.font.bold, r_f.font.color.rgb = "Courier New", True, RGBColor(31, 78, 120)

    doc_h(doc, "4. Exploratory Data Analysis Insights", 1)
    for finding in [
        "CIBIL Score: Applications with score > 650 have >92% approval rate — strongest single predictor.",
        "Asset Coverage Ratio: total_assets / loan_amount > 1.25x correlates with 95%+ approval probability.",
        "Income-to-Loan Ratio: Higher income relative to loan amount is a strong positive signal.",
        "Class Balance: 62.2% Approved vs 37.8% Rejected — evaluated by Precision, Recall, F1.",
    ]:
        doc.add_paragraph(finding, style="List Bullet")

    doc_h(doc, "5. Data Preprocessing Pipeline", 1)
    for step in [
        "Removed loan_id, education, self_employed (non-predictive or low importance).",
        "Aggregated 4 asset columns into total_assets_value.",
        "Target encoding: loan_status → binary integer (Approved=1, Rejected=0).",
        "StandardScaler applied to all 6 final predictor features.",
        "Train/Test Split: 80% training (3,415) / 20% testing (854), random_state=42.",
    ]:
        doc.add_paragraph(step, style="List Bullet")

    doc_h(doc, "6. Model Training & Benchmark Evaluation", 1)
    bench = doc.add_table(4, 5); bench.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ci, h in enumerate(["Model Algorithm", "Accuracy", "Precision", "Recall", "F1-Score"]):
        c = bench.rows[0].cells[ci]; c.text = h
        shade_cell(c, "1F4E78")
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    model_data = [
        ("Logistic Regression",              "90.75%", "92.24%", "93.10%", "92.66%"),
        ("Decision Tree Classifier",         "97.89%", "97.26%", "99.44%", "98.34%"),
        ("Random Forest Classifier (Chosen)","97.89%", "97.61%", "99.07%", "98.33%"),
    ]
    for ri, row in enumerate(model_data, 1):
        for ci, val in enumerate(row):
            c = bench.rows[ri].cells[ci]; c.text = val
            cell_margins(c, 80, 80, 100, 100)
            if ri == 3:
                shade_cell(c, "EBF3FB"); c.paragraphs[0].runs[0].font.bold = True

    doc_h(doc, "7. Fullstack Deployment Architecture", 1)
    doc_h(doc, "7.1 FastAPI REST Backend (backend/app.py)", 2)
    doc.add_paragraph("Endpoint: POST http://127.0.0.1:8000/predict")
    doc.add_paragraph("Input (JSON): no_of_dependents, income_annum, loan_amount, loan_term, cibil_score, total_assets_value")
    doc.add_paragraph("Response (JSON): prediction, approval_probability (e.g. 0.6109), rejection_probability, probability_percent ('61.09%')")
    doc_h(doc, "7.2 Streamlit Interactive UI (frontend/streamlit_app.py)", 2)
    for item in [
        "Slider/number inputs for all 6 predictor features.",
        "Real-time animated progress bar showing approval probability.",
        "Color-coded decision card: Approved (green) or Rejected (red).",
        "Run: streamlit run frontend/streamlit_app.py",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc_h(doc, "8. Business Recommendations", 1)
    for rec in [
        "Tier 1 (Auto-Approve): probability >= 80% — skip manual review.",
        "Tier 2 (Manual Underwriting): probability 40%–79% — route to credit officer.",
        "Tier 3 (Auto-Reject): probability < 40% — instant rejection.",
        "CIBIL Threshold Policy: Flag all applications with cibil_score < 550 for enhanced due diligence.",
        "Quarterly Model Retraining: Retrain on new loan book data every quarter to prevent concept drift.",
    ]:
        doc.add_paragraph(rec, style="List Number")

    doc.save(DOCX_OUT)
    print(f"Saved: {DOCX_OUT}")


def build_pptx():
    print("Generating Loan_Approval_Presentation.pptx ...")
    prs = Presentation()
    prs.slide_width, prs.slide_height = PI(13.333), PI(7.5)
    blank = prs.slide_layouts[6]
    NAVY  = PRC(31, 78, 120); DARK  = PRC(15, 32, 67); WHITE = PRC(255, 255, 255)
    GRAY  = PRC(89, 89, 89);  LGRAY = PRC(242, 244, 247); GOLD  = PRC(230, 150, 0)

    def header(slide, title, subtitle="PROJECT 1: LOAN APPROVAL PREDICTION SYSTEM"):
        tb = slide.shapes.add_textbox(PI(0.8), PI(0.25), PI(11.7), PI(1.1))
        tf = tb.text_frame; tf.word_wrap = True
        p1 = tf.paragraphs[0]; p1.text = subtitle
        p1.font.name, p1.font.size, p1.font.bold, p1.font.color.rgb = "Arial", PPt(9), True, GOLD
        p2 = tf.add_paragraph(); p2.text = title
        p2.font.name, p2.font.size, p2.font.bold, p2.font.color.rgb = "Arial", PPt(24), True, NAVY

    # SLIDE 1 — Title
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PI(13.333), PI(7.5))
    bg.fill.solid(); bg.fill.fore_color.rgb = DARK; bg.line.fill.background()
    tb = s.shapes.add_textbox(PI(1.0), PI(1.8), PI(11.3), PI(4.0))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "LOAN APPROVAL PREDICTION SYSTEM"
    p.font.name, p.font.size, p.font.bold, p.font.color.rgb = "Arial", PPt(34), True, WHITE
    p2 = tf.add_paragraph()
    p2.text = "Automated Credit Risk Classification  ·  FastAPI REST  ·  Streamlit Web UI"
    p2.font.name, p2.font.size, p2.font.color.rgb = "Arial", PPt(16), GOLD; p2.space_before = PPt(12)
    p3 = tf.add_paragraph()
    p3.text = "Random Forest Classifier  ·  97.89% Accuracy  ·  Continuous Probability Scoring"
    p3.font.name, p3.font.size, p3.font.color.rgb = "Arial", PPt(13), PRC(190, 210, 230); p3.space_before = PPt(24)

    # SLIDE 2 — Project Overview
    s2 = prs.slides.add_slide(blank)
    header(s2, "Project Overview & Business Problem")
    for i, (h, b) in enumerate([
        ("BUSINESS PROBLEM", "Manual loan underwriting is slow, inconsistent, and exposes lenders to avoidable credit risk."),
        ("ML SOLUTION", "Random Forest trained on 4,269 records delivers real-time decisions with calibrated probability scores."),
        ("DEPLOYMENT", "FastAPI REST backend + Streamlit interactive dashboard serving predictions from any browser."),
    ]):
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(0.8 + i * 4.04), PI(1.55), PI(3.64), PI(5.2))
        card.fill.solid(); card.fill.fore_color.rgb = LGRAY; card.line.color.rgb = NAVY
        tf = card.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = h; p.font.bold, p.font.size, p.font.color.rgb = True, PPt(13), NAVY
        pb = tf.add_paragraph(); pb.text = b; pb.font.size, pb.font.color.rgb = PPt(12), GRAY; pb.space_before = PPt(14)

    # SLIDE 3 — Dataset & Feature Engineering
    s3 = prs.slides.add_slide(blank)
    header(s3, "Dataset & Feature Engineering")
    tb = s3.shapes.add_textbox(PI(0.8), PI(1.6), PI(5.8), PI(5.1))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "DATASET PROFILE"
    p.font.bold, p.font.color.rgb, p.font.size = True, NAVY, PPt(14)
    for item in ["loan_approval_dataset.csv  ·  4,269 records", "Zero missing values (clean dataset)",
                 "Class: 62.2% Approved / 37.8% Rejected", "Train 80% / Test 20% (random_state=42)"]:
        pb = tf.add_paragraph(); pb.text = "• " + item
        pb.font.size, pb.font.color.rgb = PPt(12), GRAY; pb.space_before = PPt(8)
    tb2 = s3.shapes.add_textbox(PI(6.8), PI(1.6), PI(5.7), PI(5.1))
    tf2 = tb2.text_frame; tf2.word_wrap = True
    p = tf2.paragraphs[0]; p.text = "ASSET AGGREGATION FORMULA"
    p.font.bold, p.font.color.rgb, p.font.size = True, NAVY, PPt(14)
    p_code = tf2.add_paragraph()
    p_code.text = "total_assets_value =\nresidential + commercial\n+ luxury + bank_asset_value"
    p_code.font.name, p_code.font.bold, p_code.font.size = "Courier New", True, PPt(11)
    p_code.font.color.rgb = NAVY; p_code.space_before = PPt(12)
    for item in ["6 final predictor features used", "StandardScaler normalization applied"]:
        pb = tf2.add_paragraph(); pb.text = "• " + item
        pb.font.size, pb.font.color.rgb = PPt(12), GRAY; pb.space_before = PPt(10)

    # SLIDE 4 — EDA
    s4 = prs.slides.add_slide(blank)
    header(s4, "Exploratory Data Analysis — Key Insights")
    for i, (h, b) in enumerate([
        ("CIBIL SCORE THRESHOLD", "Applications with CIBIL > 650 carry >92% approval rate — single strongest predictor."),
        ("ASSET COVERAGE RATIO", "total_assets / loan_amount > 1.25x correlates with 95%+ approval probability."),
        ("INCOME-TO-LOAN RATIO", "Higher income relative to loan amount is a strong positive approval signal."),
        ("CLASS BALANCE", "62.2% Approved vs 37.8% Rejected — model evaluated by Precision, Recall, F1."),
    ]):
        top = PI(1.55 + i * 1.45)
        box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(0.8), top, PI(11.7), PI(1.2))
        box.fill.solid(); box.fill.fore_color.rgb = LGRAY; box.line.color.rgb = NAVY
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = h; p.font.bold, p.font.color.rgb, p.font.size = True, NAVY, PPt(13)
        pb = tf.add_paragraph(); pb.text = b; pb.font.size, pb.font.color.rgb = PPt(12), GRAY

    # SLIDE 5 — Model Benchmarking
    s5 = prs.slides.add_slide(blank)
    header(s5, "Model Benchmarking — 3-Algorithm Comparison")
    t = s5.shapes.add_table(4, 5, PI(0.8), PI(1.8), PI(11.7), PI(3.6)).table
    for i, h in enumerate(["Model Algorithm", "Accuracy", "Precision", "Recall", "F1-Score"]):
        c = t.cell(0, i); c.text = h
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        p = c.text_frame.paragraphs[0]
        p.font.bold, p.font.color.rgb, p.alignment = True, WHITE, PP_ALIGN.CENTER
    for ri, row in enumerate([
        ("Logistic Regression",       "90.75%", "92.24%", "93.10%", "92.66%"),
        ("Decision Tree Classifier",  "97.89%", "97.26%", "99.44%", "98.34%"),
        ("Random Forest (Chosen)",    "97.89%", "97.61%", "99.07%", "98.33%"),
    ], 1):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.text = val
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if ci > 0 else PP_ALIGN.LEFT
            if ri == 3:
                c.fill.solid(); c.fill.fore_color.rgb = PRC(235, 241, 245); p.font.bold = True

    # SLIDE 6 — Probability Scoring
    s6 = prs.slides.add_slide(blank)
    header(s6, "Continuous Probability Risk Scoring")
    for i, (h, b) in enumerate([
        ("WHY PROBABILITIES?", "Binary 0/1 decisions discard risk magnitude. Continuous probabilities enable tiered decisioning and portfolio risk stratification."),
        ("HOW IT WORKS", "Random Forest aggregates 100 decision tree votes. Fraction voting 'Approved' is returned as approval_probability (e.g. 0.6109 = 61.09%)."),
        ("RISK TIERS", "Tier 1: >= 80%  Auto-Approve\nTier 2: 40–79%  Manual Review\nTier 3: < 40%   Auto-Reject"),
    ]):
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(0.8 + i * 4.04), PI(1.6), PI(3.64), PI(5.0))
        card.fill.solid(); card.fill.fore_color.rgb = LGRAY; card.line.color.rgb = NAVY
        tf = card.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = h; p.font.bold, p.font.color.rgb, p.font.size = True, NAVY, PPt(13)
        pb = tf.add_paragraph(); pb.text = b; pb.font.size, pb.font.color.rgb = PPt(12), GRAY; pb.space_before = PPt(14)

    # SLIDE 7 — Fullstack Architecture
    s7 = prs.slides.add_slide(blank)
    header(s7, "Fullstack System Architecture")
    for i, (h, sub, b) in enumerate([
        ("STREAMLIT UI", "frontend/streamlit_app.py", "• 6-feature input sliders\n• Animated progress bar\n• Color-coded Approved/Rejected card"),
        ("FASTAPI BACKEND", "backend/app.py  POST /predict", "• Scales input via scaler.pkl\n• Returns decimal & % probability\n• JSON response with prediction"),
        ("RANDOM FOREST MODEL", "backend/model.pkl", "• 100-tree ensemble\n• Trained on 3,415 samples\n• 97.89% Test Accuracy"),
    ]):
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(0.8 + i * 4.04), PI(1.6), PI(3.64), PI(5.2))
        card.fill.solid(); card.fill.fore_color.rgb = LGRAY; card.line.color.rgb = NAVY
        tf = card.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = h; p.font.bold, p.font.color.rgb, p.font.size = True, NAVY, PPt(13)
        p.alignment = PP_ALIGN.CENTER
        ps = tf.add_paragraph(); ps.text = sub; ps.font.color.rgb, ps.font.size = GOLD, PPt(10); ps.alignment = PP_ALIGN.CENTER
        pb = tf.add_paragraph(); pb.text = "\n" + b; pb.font.size, pb.font.color.rgb = PPt(12), GRAY

    # SLIDE 8 — Business Recommendations
    s8 = prs.slides.add_slide(blank)
    header(s8, "Business Recommendations & Roadmap")
    for i, (h, b) in enumerate([
        ("TIERED UNDERWRITING", "Auto-approve >= 80%  ·  Manual review 40–79%  ·  Auto-reject < 40%"),
        ("CIBIL POLICY", "Flag all applications with CIBIL < 550 for enhanced credit due diligence"),
        ("COLLATERAL ENFORCEMENT", "Require asset coverage ratio > 1.25x for loan amounts above INR 10 million"),
        ("MODEL GOVERNANCE", "Quarterly retraining on new loan book data to prevent concept drift"),
    ]):
        top = PI(1.6 + i * 1.42)
        box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(0.8), top, PI(11.7), PI(1.15))
        box.fill.solid(); box.fill.fore_color.rgb = LGRAY; box.line.color.rgb = NAVY
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = h; p.font.bold, p.font.color.rgb, p.font.size = True, NAVY, PPt(13)
        pb = tf.add_paragraph(); pb.text = b; pb.font.size, pb.font.color.rgb = PPt(12), GRAY

    prs.save(PPTX_OUT)
    print(f"Saved: {PPTX_OUT}")


if __name__ == "__main__":
    build_docx()
    build_pptx()
    print("\nProject 1 documents generated successfully.")
